# -*- coding: utf-8 -*-
"""Reading and writing the attribute workbook.

Non-GIS colleagues edit attributes in Excel. Getting the edits back currently
means a manual join and nobody sees what changed. This module is the file half
of making that round trip safe; the diffing lives in the algorithms.

**OGR's XLSX driver is the dependency-free core** for both directions, so the
tool works on a bare QGIS install. **openpyxl is an optional enhancement**: a
frozen header row, a locked and shaded key column, sensible column widths, and
dropdown validation on domain-backed fields. When it is absent the workbook is
plain and `feedback` says so — there is no reliable way to declare a Python
dependency in a QGIS plugin, and a hard import at start-up would take the whole
provider down with it.

The workbook has two sheets. `data` holds the key column first, then the chosen
fields, with the field names exactly as the layer spells them — not display
names, because the import side matches on them. `_kga_meta` is a hidden sheet
holding one JSON blob describing where the data came from, so the import can
tell a sheet that belongs to this layer from one that does not.
"""

import datetime
import json
import os

from .compat import T_DATE, T_DATETIME, T_STRING, T_TIME, type_key
from .schema import is_null

DATA_SHEET = 'data'
META_SHEET = '_kga_meta'

#: Where the export records what it just wrote, so the import can offer it
#: without the user hunting through a file dialog. In QgsSettings rather than a
#: module global because the round trip is not one session: you export, close
#: QGIS, edit the sheet over two days, then come back to import it.
SETTINGS_LAST_EXPORT = 'KGA/xlsx/last_export'

#: Bumped when the sheet layout changes in a way an old reader cannot handle.
WORKBOOK_FORMAT = 1

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    HAS_OPENPYXL = True
except ImportError:                         # pragma: no cover - bare install
    openpyxl = None
    HAS_OPENPYXL = False

#: Excel refuses a formula-looking cell and silently mangles a long validation
#: list; 255 characters is the documented limit for an inline list.
_VALIDATION_LIMIT = 255

KEY_FILL = 'FFF2E8CF'
HEADER_FILL = 'FF10333B'


class XlsxError(Exception):
    """Anything that stops a workbook being written or read, phrased for a user."""


# ------------------------------------------------------------------- writing

def write_workbook(path, layer, field_names, key_field, features,
                   value_maps=None, feedback=None):
    """Write the `data` and `_kga_meta` sheets.

    `features` is an iterable of `QgsFeature`; `value_maps` maps a field name to
    a list of allowed codes, which become dropdowns when openpyxl is present.

    Returns the number of rows written.
    """
    if not HAS_OPENPYXL:
        return _write_with_ogr(path, layer, field_names, key_field, features,
                               feedback)

    fields = layer.fields()
    columns = _ordered_columns(field_names, key_field)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = DATA_SHEET

    header_font = Font(bold=True, color='FFFFFFFF')
    header_fill = PatternFill('solid', fgColor=HEADER_FILL)
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center')

    key_fill = PatternFill('solid', fgColor=KEY_FILL)
    indexes = [fields.indexOf(name) for name in columns]
    rows = 0

    for feature in features:
        rows += 1
        for column, field_index in enumerate(indexes, start=1):
            value = None if field_index < 0 else feature.attribute(field_index)
            cell = sheet.cell(row=rows + 1, column=column,
                              value=_excel_value(value))
            if columns[column - 1] == key_field:
                # Shaded and locked: the key is what the import matches on, so
                # an edited key silently becomes a row that matches nothing.
                cell.fill = key_fill
                cell.protection = Protection(locked=True)
            else:
                cell.protection = Protection(locked=False)

    sheet.freeze_panes = 'A2'
    for column, name in enumerate(columns, start=1):
        width = min(40, max(10, len(name) + 4))
        sheet.column_dimensions[get_column_letter(column)].width = width

    _add_validations(sheet, columns, value_maps or {}, rows, feedback)

    # Sheet protection is what makes the locked key column actually stick.
    # No password: this is a guard rail, not a security boundary, and a
    # password a colleague does not have would just get the sheet copied.
    sheet.protection.sheet = True
    sheet.protection.selectLockedCells = False

    meta_sheet = book.create_sheet(META_SHEET)
    meta_sheet['A1'] = json.dumps(
        build_metadata(layer, columns, key_field, rows), ensure_ascii=False)
    meta_sheet.sheet_state = 'hidden'

    _ensure_directory(path)
    try:
        book.save(path)
    except OSError as exc:
        raise XlsxError('Could not write {}: {}'.format(path, exc))
    return rows


def _add_validations(sheet, columns, value_maps, rows, feedback):
    """Dropdowns on domain-backed columns.

    This is the payoff for building the Domain Manager first: a domain defined
    once becomes a dropdown in the sheet the surveyor fills in, and the data
    comes back already valid.
    """
    if not rows:
        return
    for column, name in enumerate(columns, start=1):
        codes = value_maps.get(name)
        if not codes:
            continue
        formula = '"{}"'.format(','.join(str(c).replace('"', "'")
                                         for c in codes))
        if len(formula) > _VALIDATION_LIMIT:
            _say(feedback, 'The domain on "{}" has too many values for an '
                           'Excel dropdown ({} characters); the column is left '
                           'free-text.'.format(name, len(formula)))
            continue
        validation = DataValidation(type='list', formula1=formula,
                                    allow_blank=True, showDropDown=False)
        validation.error = 'Pick one of the listed values.'
        validation.errorTitle = 'Not an allowed value'
        validation.prompt = 'Allowed: {}'.format(', '.join(str(c) for c in codes))
        sheet.add_data_validation(validation)
        letter = get_column_letter(column)
        validation.add('{0}2:{0}{1}'.format(letter, rows + 1))


def _write_with_ogr(path, layer, field_names, key_field, features, feedback):
    """The no-openpyxl path: a plain sheet through OGR's XLSX driver.

    No frozen header, no dropdowns and no metadata sheet — OGR writes tables,
    not workbooks — so the import side falls back to matching on the header row.
    """
    from osgeo import gdal, ogr

    _say(feedback, 'openpyxl is not installed, so the workbook is written '
                   'plain: no frozen header, no locked key column and no '
                   'dropdowns. The round trip still works.')

    columns = _ordered_columns(field_names, key_field)
    fields = layer.fields()
    _ensure_directory(path)

    driver = ogr.GetDriverByName('XLSX')
    if driver is None:
        raise XlsxError(
            'This GDAL build has no XLSX driver, and openpyxl is not '
            'installed either, so no workbook can be written.')

    if os.path.exists(path):
        try:
            os.remove(path)
        except OSError as exc:
            raise XlsxError('Could not replace {}: {}'.format(path, exc))

    gdal.PushErrorHandler('CPLQuietErrorHandler')
    try:
        dataset = driver.CreateDataSource(path)
        if dataset is None:
            raise XlsxError('Could not create {}.'.format(path))
        sheet = dataset.CreateLayer(DATA_SHEET, None, ogr.wkbNone)
        for name in columns:
            sheet.CreateField(ogr.FieldDefn(name, ogr.OFTString))

        definition = sheet.GetLayerDefn()
        indexes = [fields.indexOf(name) for name in columns]
        rows = 0
        for feature in features:
            rows += 1
            out = ogr.Feature(definition)
            for column, field_index in enumerate(indexes):
                value = None if field_index < 0 else feature.attribute(field_index)
                if not is_null(value):
                    out.SetField(column, _as_cell_text(value))
            sheet.CreateFeature(out)
            out = None
        dataset = None
        return rows
    finally:
        gdal.PopErrorHandler()


def remember_export(path, layer, key_field, row_count=0):
    """Record the workbook just written, for the import side to offer back."""
    from qgis.core import QgsSettings
    payload = {
        'path': os.path.abspath(path),
        'layer_id': layer.id(),
        'layer_name': layer.name(),
        'key_field': key_field,
        'row_count': row_count,
        'exported': datetime.datetime.now().isoformat(timespec='seconds'),
    }
    QgsSettings().setValue(SETTINGS_LAST_EXPORT,
                           json.dumps(payload, ensure_ascii=False))
    return payload


def last_export():
    """The remembered export, or `{}`.

    Returns nothing when the file has since been moved or deleted, so a stale
    entry can never prefill a path that would only fail at run time.
    """
    from qgis.core import QgsSettings

    raw = QgsSettings().value(SETTINGS_LAST_EXPORT, '')
    if not raw:
        return {}
    try:
        payload = json.loads(raw)
    except (ValueError, TypeError):
        return {}
    if not isinstance(payload, dict):
        return {}
    path = payload.get('path')
    if not path or not os.path.exists(path):
        return {}
    return payload


def forget_export():
    """Drop the remembered export. Used by the tests, and harmless otherwise."""
    from qgis.core import QgsSettings
    QgsSettings().remove(SETTINGS_LAST_EXPORT)


def open_externally(path):
    """Hand `path` to whatever the desktop opens .xlsx with.

    ``(opened, reason)``. Refuses when there is no GUI — a Processing run from a
    model, a batch job or a head-less script must not start launching Excel —
    so the caller can say why nothing happened.
    """
    try:
        from qgis.utils import iface
    except ImportError:                     # pragma: no cover - outside QGIS
        iface = None
    if iface is None:
        return False, 'not running in the QGIS desktop'
    if not path or not os.path.exists(path):
        return False, 'the file was not written'

    from qgis.PyQt.QtCore import QUrl
    from qgis.PyQt.QtGui import QDesktopServices

    if QDesktopServices.openUrl(QUrl.fromLocalFile(os.path.abspath(path))):
        return True, ''
    return False, ('no application is associated with .xlsx on this machine')


def build_metadata(layer, columns, key_field, row_count):
    """What the import side reads to know what it is looking at."""
    from .schema import field_signature
    return {
        'format': WORKBOOK_FORMAT,
        'generator': 'KGA Toolbox',
        'exported': datetime.datetime.now().isoformat(timespec='seconds'),
        'layer_id': layer.id(),
        'layer_name': layer.name(),
        'source': layer.source(),
        'crs': layer.crs().authid(),
        'key_field': key_field,
        'columns': list(columns),
        'fields': field_signature(layer.fields()),
        'row_count': row_count,
    }


# ------------------------------------------------------------------- reading

def read_workbook(path):
    """``(rows, metadata)`` where `rows` is a list of ``{column: value}``.

    Values come back as Python types where the file records them — openpyxl
    hands back real dates and numbers — and as text otherwise. Turning them
    into what the target field wants is `schema.coerce`'s job, not this one's.
    """
    if not os.path.exists(path):
        raise XlsxError('{} does not exist.'.format(path))
    if HAS_OPENPYXL:
        return _read_with_openpyxl(path)
    return _read_with_ogr(path)


def _read_with_openpyxl(path):
    try:
        book = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise XlsxError('Could not read {}: {}'.format(path, exc))

    try:
        sheet = book[DATA_SHEET] if DATA_SHEET in book.sheetnames \
            else book.worksheets[0]
        rows = _rows_from_openpyxl(sheet)
        metadata = {}
        if META_SHEET in book.sheetnames:
            raw = book[META_SHEET]['A1'].value
            if raw:
                try:
                    metadata = json.loads(raw)
                except ValueError:
                    metadata = {}
        return rows, metadata
    finally:
        book.close()


def _rows_from_openpyxl(sheet):
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration:
        return []
    columns = [_text(name) for name in header]
    if not any(columns):
        return []

    rows = []
    for values in iterator:
        if all(is_null(value) for value in values):
            continue                        # trailing blank rows Excel leaves
        row = {}
        for index, name in enumerate(columns):
            if not name:
                continue
            row[name] = values[index] if index < len(values) else None
        rows.append(row)
    return rows


def _read_with_ogr(path):
    from osgeo import gdal

    gdal.PushErrorHandler('CPLQuietErrorHandler')
    try:
        try:
            # HEADERS=FORCE is not optional. The driver's automatic detection
            # gives up when every column is text — which is exactly what an
            # attribute sheet looks like — and then treats the header row as
            # data and names the columns Field1, Field2, … So the header would
            # silently arrive as a bogus extra row and no column would match a
            # field name.
            dataset = gdal.OpenEx(path, gdal.OF_VECTOR,
                                  open_options=['HEADERS=FORCE'])
        except Exception as exc:
            raise XlsxError('Could not read {}: {}'.format(path, exc))
        if dataset is None:
            raise XlsxError('Could not read {} as a workbook.'.format(path))

        sheet = dataset.GetLayerByName(DATA_SHEET)
        if sheet is None:
            sheet = dataset.GetLayer(0)
        if sheet is None:
            raise XlsxError('{} has no sheets.'.format(path))

        definition = sheet.GetLayerDefn()
        columns = [definition.GetFieldDefn(i).GetName()
                   for i in range(definition.GetFieldCount())]

        rows = []
        for feature in sheet:
            row = {name: feature.GetField(index)
                   for index, name in enumerate(columns)}
            if all(is_null(v) for v in row.values()):
                continue
            rows.append(row)

        metadata = {}
        meta_layer = dataset.GetLayerByName(META_SHEET)
        if meta_layer is not None:
            meta_layer.ResetReading()
            first = meta_layer.GetNextFeature()
            if first is not None and first.GetFieldCount():
                try:
                    metadata = json.loads(first.GetField(0) or '{}')
                except ValueError:
                    metadata = {}
        dataset = None
        return rows, metadata
    finally:
        gdal.PopErrorHandler()


# ------------------------------------------------------------------ comparing

def describe_mismatch(metadata, layer):
    """Reasons this workbook may not belong to this layer, as a list of strings.

    Deliberately advisory: applying a sheet to a copy of the layer it came from
    is legitimate and common. The caller warns and keeps the dry run on rather
    than refusing.
    """
    problems = []
    if not metadata:
        return ['The workbook has no KGA metadata sheet, so there is no way to '
                'confirm it came from this layer. Check the report carefully.']

    if metadata.get('layer_id') and metadata['layer_id'] != layer.id():
        problems.append(
            'The workbook was exported from layer "{}", not from "{}".'.format(
                metadata.get('layer_name', metadata['layer_id']), layer.name()))

    recorded = {f['name']: f for f in metadata.get('fields') or []}
    if recorded:
        have = {f.name(): f for f in layer.fields()}
        missing = [n for n in recorded if n not in have]
        if missing:
            problems.append(
                'The layer no longer has field(s) the workbook records: '
                '{}.'.format(', '.join(sorted(missing))))
        for name, spec in recorded.items():
            field = have.get(name)
            if field is None:
                continue
            if type_key(field.type()) != spec.get('type'):
                problems.append(
                    'Field "{}" was {} when exported and is {} now.'.format(
                        name, spec.get('type'), type_key(field.type())))
    return problems


# ------------------------------------------------------------------- helpers

def _ordered_columns(field_names, key_field):
    """Key column first, then the chosen fields, each named once."""
    columns = []
    if key_field:
        columns.append(key_field)
    for name in field_names:
        if name not in columns:
            columns.append(name)
    return columns


def _excel_value(value):
    """A value openpyxl can write without complaint."""
    from qgis.PyQt.QtCore import QDate, QDateTime, QTime, QVariant

    if isinstance(value, QVariant):
        value = None if value.isNull() else value.value()
    if value is None:
        return None
    if isinstance(value, QDateTime):
        return value.toPyDateTime()
    if isinstance(value, QDate):
        return datetime.date(value.year(), value.month(), value.day())
    if isinstance(value, QTime):
        return datetime.time(value.hour(), value.minute(), value.second())
    if isinstance(value, (int, float, bool, str,
                          datetime.date, datetime.datetime, datetime.time)):
        return value
    if isinstance(value, (bytes, bytearray)):
        return '<binary, {} bytes>'.format(len(value))
    return str(value)


def _as_cell_text(value):
    prepared = _excel_value(value)
    if prepared is None:
        return ''
    if isinstance(prepared, (datetime.date, datetime.datetime, datetime.time)):
        return prepared.isoformat()
    return str(prepared)


def _text(value):
    return '' if value is None else str(value).strip()


def _ensure_directory(path):
    directory = os.path.dirname(path)
    if directory and not os.path.isdir(directory):
        os.makedirs(directory)


def _say(feedback, message):
    if feedback is not None:
        feedback.pushInfo(message)


def value_maps_for(layer, field_names):
    """`{field: [code, ...]}` from each field's ValueMap editor widget.

    Reads the widget rather than the domain directly, so a value map a user set
    by hand works exactly like one a field domain produced.
    """
    maps = {}
    fields = layer.fields()
    for name in field_names:
        index = fields.indexOf(name)
        if index < 0:
            continue
        setup = layer.editorWidgetSetup(index)
        if setup.type() != 'ValueMap':
            continue
        codes = []
        for entry in setup.config().get('map') or []:
            if isinstance(entry, dict):
                codes.extend(str(code) for code in entry.values())
            else:
                codes.append(str(entry))
        if codes:
            maps[name] = codes
    return maps
